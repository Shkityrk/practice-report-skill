"""Библиотека сборки отчёта по практической работе в .docx (оформление по образцу отчётов РТУ МИРЭА; шапку и реквизиты можно менять в config.json).

Параметры оформления: Times New Roman 14, интервал 1.5, выравнивание по ширине, абзацный отступ 1.25 см,
поля 3 / 1.5 / 2 / 2 см (лево/право/верх/низ), номер страницы снизу по центру (на титульном листе нет),
заголовки разделов – по центру, жирные, ПРОПИСНЫЕ, каждый с новой страницы; рисунки и таблицы с подписями
«Рисунок N – …» (под рисунком) и «Таблица N – …» (над таблицей); листинги – Courier New 10, серый фон.

Использование (см. templates/build_report.py):
    from report_lib import Report
    r = Report(cfg)            # cfg – dict из config.json
    r.title_page()
    r.toc()
    r.heading("ЦЕЛЬ РАБОТЫ"); r.para("...")
    r.figure("01-x.png", "Подпись")
    r.save()
"""
import datetime
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
MONO = "Courier New"
J, L, C = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER


def _field(run, instr):
    """Вставляет в run поле Word (PAGE, TOC ...)."""
    for kind, val in (("begin", None), ("instr", instr), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = val
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


class Report:
    def __init__(self, cfg, shots_dir=None, out_path=None):
        self.cfg = cfg
        self.shots = shots_dir or cfg.get("shots_dir", "screenshots")
        self.out = out_path or cfg["out"]
        self.fig_n = 0
        self.tab_n = 0
        d = self.doc = Document()

        st = d.styles["Normal"]
        st.font.name, st.font.size = FONT, Pt(14)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.space_after = Pt(0)

        h1 = d.styles["Heading 1"]
        h1.font.name, h1.font.size, h1.font.bold = FONT, Pt(14), True
        h1.font.color.rgb = RGBColor(0, 0, 0)
        for a in ("w:eastAsia", "w:ascii", "w:hAnsi"):
            h1.element.rPr.rFonts.set(qn(a), FONT)
        pf = h1.paragraph_format
        pf.alignment, pf.space_before, pf.space_after = C, Pt(0), Pt(12)
        pf.line_spacing, pf.keep_with_next = 1.5, True

        sec = d.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin, sec.right_margin = Cm(3.0), Cm(1.5)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        sec.different_first_page_header_footer = True  # на титульном листе номера нет
        fp = sec.footer.paragraphs[0]
        fp.alignment = C
        run = fp.add_run()
        _field(run, "PAGE")
        run.font.name, run.font.size = FONT, Pt(14)

    # ---------- базовые элементы ----------
    def para(self, text="", align=J, indent=True, bold=False, italic=False, size=None,
             spacing=None, keep_next=False, space_after=None):
        p = self.doc.add_paragraph()
        p.alignment = align
        if indent:
            p.paragraph_format.first_line_indent = Cm(1.25)
        if spacing:
            p.paragraph_format.line_spacing = spacing
        if keep_next:
            p.paragraph_format.keep_with_next = True
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        if text:
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
            if size:
                r.font.size = Pt(size)
        return p

    def sub(self, text):
        """Подзаголовок (жирный, по левому краю, без отступа), напр. «1 Подготовка окружения»."""
        return self.para(text, align=L, indent=False, bold=True, keep_next=True)

    def heading(self, text, page_break=True):
        h = self.doc.add_heading(text, level=1)
        if page_break:
            h.paragraph_format.page_break_before = True
        return h

    def code(self, text):
        """Листинг: Courier New 10, одинарный интервал, серый фон. Для кода/YAML в тексте отчёта."""
        lines = text.strip("\n").split("\n")
        for i, line in enumerate(lines):
            p = self.doc.add_paragraph()
            p.alignment = L
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.left_indent = Cm(0.3)
            p.paragraph_format.keep_together = True
            if i < len(lines) - 1:
                p.paragraph_format.keep_with_next = True
            ppr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "F2F2F2")
            ppr.append(shd)
            r = p.add_run(line if line else " ")
            r.font.name, r.font.size = MONO, Pt(10)
            r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), MONO)
        self.para(spacing=1.0, indent=False)

    def figure(self, fname, caption, width_cm=16.0):
        """Рисунок по центру и подпись «Рисунок N – …» под ним. Возвращает номер рисунка."""
        path = fname if os.path.isabs(fname) else os.path.join(self.shots, fname)
        p = self.doc.add_paragraph()
        p.alignment = C
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Cm(width_cm))
        self.fig_n += 1
        c = self.doc.add_paragraph()
        c.alignment = C
        c.paragraph_format.line_spacing = 1.0
        c.paragraph_format.space_after = Pt(12)
        c.add_run(f"Рисунок {self.fig_n} – {caption}")
        return self.fig_n

    def bullets(self, items):
        """Список через тире. Элемент – строка или кортеж (жирное_начало, продолжение)."""
        for it in items:
            p = self.para(indent=False)
            p.paragraph_format.left_indent = Cm(1.25)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            if isinstance(it, tuple):
                p.add_run("– ")
                r = p.add_run(it[0])
                r.bold = True
                p.add_run(it[1])
            else:
                p.add_run("– " + it)

    def table(self, rows, widths_cm, caption=None):
        """Таблица с подписью «Таблица N – …» НАД таблицей; первая строка – заголовок."""
        if caption:
            self.tab_n += 1
            self.para(f"Таблица {self.tab_n} – {caption}", align=L, indent=False, keep_next=True,
                      spacing=1.0, space_after=4)
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(rows):
            for j, txt in enumerate(row):
                cell = t.cell(i, j)
                cell.width = Cm(widths_cm[j])
                cp = cell.paragraphs[0]
                cp.paragraph_format.line_spacing = 1.0
                r = cp.add_run(txt)
                r.font.size = Pt(12)
                if i == 0:
                    r.bold = True
                    cp.alignment = C
        self.para(spacing=1.0, indent=False)
        return t

    def sources(self, items):
        """Список источников с висячим отступом."""
        for i, s in enumerate(items, 1):
            p = self.para(f"{i}. {s}", align=L, indent=False)
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.first_line_indent = Cm(-0.8)

    # ---------- титульный лист и оглавление ----------
    def _center(self, text, bold=False, size=14, after=0):
        p = self.para(text, align=C, indent=False, bold=bold, size=size, spacing=1.0)
        p.paragraph_format.space_after = Pt(after)

    def _sign_row(self, left, right):
        t = self.doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.cell(0, 0).width, t.cell(0, 1).width = Cm(10.5), Cm(6.0)
        a = t.cell(0, 0).paragraphs[0]
        a.paragraph_format.line_spacing = 1.0
        a.add_run(left)
        b = t.cell(0, 1).paragraphs[0]
        b.paragraph_format.line_spacing = 1.0
        b.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        b.add_run(right)

    def _small(self, text, after=10):
        p = self.para(text, align=L, indent=False, size=11, spacing=1.0)
        p.paragraph_format.left_indent = Cm(1.2)
        p.paragraph_format.space_after = Pt(after)

    def title_page(self):
        c = self.cfg
        logo = c.get("logo")
        if logo and os.path.exists(logo):          # логотип необязателен
            lp = self.doc.add_paragraph()
            lp.alignment = C
            lp.paragraph_format.line_spacing = 1.0
            lp.add_run().add_picture(logo, height=Cm(3.0))
        # шапка вуза: список [текст, жирный]; по умолчанию – как в отчётах РТУ МИРЭА, для другого вуза задать cfg["header"]
        header = c.get("header") or [
            ["МИНОБРНАУКИ РОССИИ", True],
            ["Федеральное государственное бюджетное образовательное учреждение", False],
            ["высшего образования", False],
            ["«МИРЭА – Российский технологический университет»", True],
            ["РТУ МИРЭА", True],
        ]
        for i, (text, bold) in enumerate(header):
            self._center(text, bold=bold, after=14 if i == len(header) - 1 else 0)
        unit = [c[k] for k in ("institute", "department") if c.get(k)]   # подразделение – только если задано
        for i, text in enumerate(unit):
            self._center(text, after=60 if i == len(unit) - 1 else 0)
        self._center(f"ПРАКТИЧЕСКАЯ РАБОТА №{c['practice_no']}", bold=True, after=14)
        self._center(f"по дисциплине «{c['discipline']}»", after=60)
        self._sign_row(f"Студент группы {c['group']}", c["student_short"])
        self._small("_______________\n(подпись студента)")
        self._sign_row("Руководитель практической работы", c["supervisor"])
        self._small("_______________\n(подпись руководителя)", after=14)
        y = str(c.get("year") or datetime.date.today().year)
        self.para(f"Работа представлена   «____» ____________ {y} г.", align=L, indent=False, spacing=1.0, space_after=10)
        self.para(f"Допущен к работе   «___» ____________ {y} г.", align=L, indent=False, spacing=1.0, space_after=70)
        self._center(f"{c.get('city', 'Москва')} {y}")

    def toc(self):
        """Оглавление-поле Word (обновляется при открытии файла: нужно согласиться на обновление полей)."""
        self.heading("СОДЕРЖАНИЕ")
        r = self.doc.add_paragraph().add_run()
        for kind, val in (("begin", None), ("instr", 'TOC \\o "1-1" \\h \\z \\u'), ("separate", None),
                          ("text", "Обновите поле (ПКМ → «Обновить поле»), чтобы построить оглавление."), ("end", None)):
            if kind in ("begin", "separate", "end"):
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            elif kind == "instr":
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = val
            else:
                el = OxmlElement("w:t")
                el.text = val
            r._r.append(el)
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        self.doc.settings.element.append(uf)

    def save(self):
        self.doc.save(self.out)
        print("saved", self.out, "| рисунков:", self.fig_n, "| таблиц:", self.tab_n)
        return self.out
